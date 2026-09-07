from pathlib import Path

from langchain_core.documents import Document
import fitz
from docx import Document as DocxDocument


def load_txt(file_path: str) -> list[Document]:
    path = Path(file_path)

    text = path.read_text(encoding="utf-8")

    return [
        Document(
            page_content=text,
            metadata={
                "source": str(path),
                "file_type": "txt",
            },
        )
    ]


def load_pdf(file_path: str) -> list[Document]:
    path = Path(file_path)

    pdf = fitz.open(file_path)

    documents = []

    for page_number, page in enumerate(pdf, start=1):
        text = page.get_text()

        if text.strip():
            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": str(path),
                        "file_type": "pdf",
                        "page": page_number,
                    },
                )
            )

    pdf.close()

    return documents


def load_docx(file_path: str) -> list[Document]:
    path = Path(file_path)

    doc = DocxDocument(file_path)

    paragraphs = [
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.text.strip()
    ]

    text = "\n".join(paragraphs)

    return [
        Document(
            page_content=text,
            metadata={
                "source": str(path),
                "file_type": "docx",
            },
        )
    ]


def load_document(file_path: str) -> list[Document]:
    path = Path(file_path)

    extension = path.suffix.lower()

    if extension == ".txt":
        return load_txt(file_path)

    if extension == ".pdf":
        return load_pdf(file_path)

    if extension == ".docx":
        return load_docx(file_path)

    raise ValueError(
        f"Unsupported file type: {extension}. "
        "Supported types are .txt, .pdf and .docx."
    )